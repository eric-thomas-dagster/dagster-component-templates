"""Committed tests for GoogleDriveIngestionComponent's `inline_text_content`
binary-document extraction (PDF/.docx/.xlsx via pdfplumber/python-docx/openpyxl).

Uses real fixture bytes throughout -- a real PDF via reportlab, a real .docx
via python-docx, a real .xlsx via openpyxl -- and mocks only the Drive API
network boundary (see conftest.py). The extraction libraries themselves are
never mocked, so these are real round-trips, not fakes.
"""
import io

import dagster as dg
import pytest

from .conftest import install_fake_google_drive, load_component_module


@pytest.fixture()
def mod():
    return load_component_module()


def _make_component(mod, **overrides):
    kwargs = dict(
        asset_name="test_asset",
        credentials={"type": "service_account", "client_email": "sa@example.com"},
        download=True,
        inline_text_content=True,
    )
    kwargs.update(overrides)
    return mod.GoogleDriveIngestionComponent(**kwargs)


def _materialize(component):
    asset_def = component.build_defs(context=None).assets[0]
    return dg.materialize(
        [asset_def],
        resources={"io_manager": dg.InMemoryIOManager()},
    )


def _make_pdf_bytes(text: str) -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 700, text)
    c.save()
    return buf.getvalue()


def _make_docx_bytes(paragraphs) -> bytes:
    import docx

    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _make_xlsx_bytes(sheet_name: str, rows) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Direct unit tests on the extraction helpers (fast, no Drive/asset plumbing) ──

def test_extract_pdf_text_real_pdf(mod):
    data = _make_pdf_bytes("Hello from a real PDF fixture")
    assert "Hello from a real PDF fixture" in mod._extract_pdf_text(data)


def test_extract_docx_text_real_docx(mod):
    data = _make_docx_bytes(["First paragraph.", "Second paragraph."])
    text = mod._extract_docx_text(data)
    assert "First paragraph." in text
    assert "Second paragraph." in text


def test_extract_xlsx_text_real_xlsx(mod):
    data = _make_xlsx_bytes("Sheet1", [["name", "amount"], ["Acme Corp", "42"]])
    text = mod._extract_xlsx_text(data)
    assert "Sheet1" in text
    assert "Acme Corp" in text
    assert "42" in text


def test_extract_pdf_text_corrupt_bytes_returns_none_without_raising(mod):
    assert mod._extract_pdf_text(b"not a real pdf at all") is None


def test_extract_docx_text_corrupt_bytes_returns_none_without_raising(mod):
    assert mod._extract_docx_text(b"not a real docx at all") is None


def test_extract_xlsx_text_corrupt_bytes_returns_none_without_raising(mod):
    assert mod._extract_xlsx_text(b"not a real xlsx at all") is None


def test_extract_pdf_text_logs_warning_on_failure(mod):
    class _FakeLog:
        def __init__(self):
            self.warnings = []

        def warning(self, msg):
            self.warnings.append(msg)

    log = _FakeLog()
    result = mod._extract_pdf_text(b"garbage", log)
    assert result is None
    assert len(log.warnings) == 1


# ── End-to-end: full asset materialization dispatches by MIME type correctly ──

def test_materialize_extracts_pdf_content_by_mime_type(mod, monkeypatch):
    pdf_bytes = _make_pdf_bytes("Discovery call notes for Acme Corp")
    files = [{
        "id": "f1", "name": "notes.pdf", "mimeType": "application/pdf",
        "size": str(len(pdf_bytes)), "modifiedTime": "2026-01-01T00:00:00Z",
    }]
    install_fake_google_drive(monkeypatch, files, {"f1": pdf_bytes})

    component = _make_component(mod)
    result = _materialize(component)
    assert result.success
    df = result.output_for_node("test_asset")
    assert "Discovery call notes for Acme Corp" in df.iloc[0]["content"]


def test_materialize_extracts_docx_content_by_mime_type(mod, monkeypatch):
    docx_bytes = _make_docx_bytes(["AE handoff notes.", "Stack: Snowflake, dbt, Airflow."])
    files = [{
        "id": "f2", "name": "handoff.docx",
        "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "size": str(len(docx_bytes)), "modifiedTime": "2026-01-01T00:00:00Z",
    }]
    install_fake_google_drive(monkeypatch, files, {"f2": docx_bytes})

    component = _make_component(mod)
    result = _materialize(component)
    assert result.success
    df = result.output_for_node("test_asset")
    assert "Stack: Snowflake, dbt, Airflow." in df.iloc[0]["content"]


def test_materialize_extracts_xlsx_content_by_mime_type(mod, monkeypatch):
    xlsx_bytes = _make_xlsx_bytes("Pricing", [["tier", "price"], ["enterprise", "99999"]])
    files = [{
        "id": "f3", "name": "pricing.xlsx",
        "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "size": str(len(xlsx_bytes)), "modifiedTime": "2026-01-01T00:00:00Z",
    }]
    install_fake_google_drive(monkeypatch, files, {"f3": xlsx_bytes})

    component = _make_component(mod)
    result = _materialize(component)
    assert result.success
    df = result.output_for_node("test_asset")
    assert "enterprise" in df.iloc[0]["content"]


def test_materialize_unparseable_pdf_bytes_falls_back_to_none_content(mod, monkeypatch):
    files = [{
        "id": "f4", "name": "broken.pdf", "mimeType": "application/pdf",
        "size": "7", "modifiedTime": "2026-01-01T00:00:00Z",
    }]
    install_fake_google_drive(monkeypatch, files, {"f4": b"garbage"})

    component = _make_component(mod)
    result = _materialize(component)
    assert result.success
    df = result.output_for_node("test_asset")
    assert df.iloc[0]["content"] is None


def test_materialize_unknown_binary_mime_type_stays_none(mod, monkeypatch):
    files = [{
        "id": "f5", "name": "image.png", "mimeType": "image/png",
        "size": "3", "modifiedTime": "2026-01-01T00:00:00Z",
    }]
    install_fake_google_drive(monkeypatch, files, {"f5": b"\x89PN"})

    component = _make_component(mod)
    result = _materialize(component)
    assert result.success
    df = result.output_for_node("test_asset")
    assert df.iloc[0]["content"] is None


def test_materialize_without_inline_text_content_never_extracts(mod, monkeypatch):
    pdf_bytes = _make_pdf_bytes("Should not be extracted")
    files = [{
        "id": "f6", "name": "notes.pdf", "mimeType": "application/pdf",
        "size": str(len(pdf_bytes)), "modifiedTime": "2026-01-01T00:00:00Z",
    }]
    install_fake_google_drive(monkeypatch, files, {"f6": pdf_bytes})

    component = _make_component(mod, inline_text_content=False)
    result = _materialize(component)
    assert result.success
    df = result.output_for_node("test_asset")
    assert "content" not in df.columns
